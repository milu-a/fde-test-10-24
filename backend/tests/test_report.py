# -*- coding: utf-8 -*-
"""M2 台账生成器回归：生成的 xlsx 重开校验——表头与模板逐字一致、关键数字与状态正确。"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook

from engine import load_all
from engine.money import D
from report import SHEET_ORDER, headers_for, write_ledger

SUPS4 = ["甲", "乙", "丙", "丁"]

QUALIFIERS = [
    {"supplier": "甲", "category": "营业执照摘要", "doc": "MAT-A-001", "submitted": "已提交",
     "validity": "有效", "note": "", "review": "待评审复核"},
    {"supplier": "乙", "category": "检测资质摘要", "doc": "", "submitted": "未提交",
     "validity": "", "note": "检测资质摘要待补充", "review": "缺件：未提交，阻塞评审"},
    {"supplier": "丁", "category": "检测资质摘要", "doc": "MAT-D-002", "submitted": "已提交",
     "validity": "2026-08-31", "note": "资质2026-08-31到期，续期材料未提供", "review": "疑点：有效性待评审确认"},
]
RATE_RAW = {"甲": "6%", "乙": "未声明", "丙": "6%", "丁": "6%"}
RATE_OK = {s: "6%" for s in ("甲", "乙", "丙", "丁")}


def _find_row(ws, col, value, start=2):
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == value:
            return r
    return None


class ReportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        ns = load_all()
        b = ns.builder
        ch = ns.events["change"]
        cls.catalog = ns.catalog
        cls.v1 = b.build("V1-原始")
        cls.v1c = b.build("V1-澄清后", amendments=ns.events["clarify"])
        cls.v2 = b.build("V2-变更后", amendments=ns.events["clarify"],
                         qty_overrides=ch["qty_overrides"], qty_basis=ch["id"],
                         qty_at=ch["at"], qty_operator=ch["operator"])
        catalog_qty = {i: c["qty"] for i, c in ns.catalog.items()}
        invoice_all = {s: {"declared": True, "invoice_type": None} for s in SUPS4}
        cls.paths = {}
        for name, led, rate, qmap, prev in (
            ("raw", cls.v1, RATE_RAW, catalog_qty, None),
            ("v1c", cls.v1c, RATE_OK, catalog_qty, None),
            ("v2", cls.v2, RATE_OK, {**catalog_qty, "C08": 72},
             {s: cls.v1c.totals[s]["hs"] for s in SUPS4}),
        ):
            cls.paths[name] = write_ledger(
                led, ns.catalog, Path(cls.tmp.name) / f"{name}.xlsx", rate,
                qualifiers=QUALIFIERS, qty_map=qmap, prev_totals=prev,
                version_note=f"台账版本 {name}",
                issue_source=cls.v1,  # 问题清单自 V1 原始累计，状态对当前版本求值
                invoice_status=invoice_all)
        cls.wb_raw = load_workbook(cls.paths["raw"])
        cls.wb_v2 = load_workbook(cls.paths["v2"])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    # ---------------- 模板保真 ----------------
    def test_sheets_and_headers_exact(self):
        self.assertEqual(self.wb_raw.sheetnames, SHEET_ORDER)
        for name in SHEET_ORDER:
            ws = self.wb_raw[name]
            got = [ws.cell(row=1, column=i).value for i in range(1, len(headers_for(name, SUPS4)) + 1)]
            self.assertEqual(got, headers_for(name, SUPS4), f"{name} 表头与模板不一致")

    # ---------------- 对比表 ----------------
    def test_db_raw_key_cells(self):
        ws = self.wb_raw["对比表"]
        r_c01 = _find_row(ws, 1, "C01")
        self.assertEqual(ws.cell(row=r_c01, column=5).value, 2289.6)      # 甲含税
        self.assertEqual(ws.cell(row=r_c01, column=6).value, "税率未给")  # 乙硬规则
        r_c07 = _find_row(ws, 1, "C07")
        self.assertEqual(ws.cell(row=r_c07, column=6).value, "税率未给")  # 乙含税列：税率未给（重复在差异列与未税合计体现）
        r_c09 = _find_row(ws, 1, "C09")
        self.assertEqual(ws.cell(row=r_c09, column=6).value, "未报")
        r_sum = _find_row(ws, 1, "合计")
        self.assertEqual(ws.cell(row=r_sum, column=5).value, 40545.0)
        self.assertEqual(ws.cell(row=r_sum, column=6).value, "—")  # 乙合计不可折算

    def test_db_v2_change_row(self):
        ws = self.wb_v2["对比表"]
        r_c08 = _find_row(ws, 1, "C08")
        self.assertEqual(ws.cell(row=r_c08, column=4).value, 72)          # 统一数量
        self.assertEqual(ws.cell(row=r_c08, column=5).value, 6487.2)      # 甲 72×90.1
        r_diff = _find_row(ws, 1, "较V1差异")
        self.assertEqual(ws.cell(row=r_diff, column=5).value, "+1,081.20")  # 6487.20-5406.00

    # ---------------- 报价明细 ----------------
    def test_mx_raw_rows_and_bing_c04(self):
        ws = self.wb_raw["报价明细"]
        # 甲15 + 乙16(含C07×2与C11占位) + 丙15(含C13占位) + 丁15 = 61 行
        self.assertEqual(ws.max_row - 1, 61)
        r = None
        for row in range(2, ws.max_row + 1):
            if (ws.cell(row=row, column=1).value == "丙"
                    and ws.cell(row=row, column=2).value == "C04"):
                r = row
        self.assertEqual(ws.cell(row=r, column=8).value, 2181.8)   # 原金额（列示）
        self.assertEqual(ws.cell(row=r, column=11).value, 2039.43) # 未税
        self.assertEqual(ws.cell(row=r, column=13).value, 2161.8)  # 含税按乘积
        self.assertEqual(ws.cell(row=r, column=7).value, "6%")

    def test_mx_yi_rate_missing_amounts(self):
        ws = self.wb_raw["报价明细"]
        for row in range(2, ws.max_row + 1):
            if (ws.cell(row=row, column=1).value == "乙"
                    and ws.cell(row=row, column=2).value == "C01"):
                self.assertEqual(ws.cell(row=row, column=11).value, 2116.8)  # 未税
                self.assertEqual(ws.cell(row=row, column=12).value, "税率未给")
                self.assertEqual(ws.cell(row=row, column=13).value, "税率未给")
                return
        self.fail("未找到乙C01明细行")

    # ---------------- 操作记录 ----------------
    def test_audit_rows(self):
        self.assertEqual(self.wb_raw["操作记录"].max_row - 1, 3)
        self.assertEqual(self.wb_v2["操作记录"].max_row - 1, 15)

    # ---------------- 需补充事项 ----------------
    def _need_rows(self, wb):
        ws = wb["需补充事项"]
        out = []
        for r in range(2, ws.max_row + 1):
            out.append(tuple(ws.cell(row=r, column=c).value for c in (2, 3, 6)))
        return out

    def test_qb_raw_open_clarify_closed_in_v2(self):
        raw = self._need_rows(self.wb_raw)
        v2 = self._need_rows(self.wb_v2)
        # 原始版：乙税率（按供应商折叠）、乙C11漏报、丙C04抄写、丁C14单位 全部待补充
        for key in (("乙", "—", "待补充"), ("乙", "C11", "待补充"),
                    ("丙", "C04", "待补充"), ("丁", "C14", "待补充")):
            self.assertIn(key, raw, f"原始版缺待补充事项 {key}")
        # V2：澄清类已关闭；乙C08区间异常不因工程量变更被关闭；乙/丁资质事项保持开放
        self.assertIn(("乙", "—", "已澄清关闭"), v2)
        self.assertIn(("丙", "C04", "已澄清关闭"), v2)
        self.assertIn(("丁", "C14", "已澄清关闭"), v2)
        self.assertIn(("乙", "C08", "待补充"), v2)
        # M5：四家税率均已声明但发票类型未说明 → 各立一项且不自动关闭
        for sup in ("甲", "乙", "丙", "丁"):
            self.assertIn((sup, "税率", "待补充"), v2)
        opens = [k for k in v2 if k[2] == "待补充"]
        self.assertIn(("乙", "资质", "待补充"), opens)
        self.assertIn(("丁", "资质", "待补充"), opens)

    def test_zg_rows(self):
        ws = self.wb_raw["资格材料清单"]
        self.assertEqual(ws.max_row - 1, len(QUALIFIERS))
        r = _find_row(ws, 1, "乙")
        self.assertEqual(ws.cell(row=r, column=3).value, "—")      # 文件编号空
        self.assertEqual(ws.cell(row=r, column=4).value, "未提交")  # 提交情况
        self.assertIn("缺件", ws.cell(row=r, column=7).value)


if __name__ == "__main__":
    unittest.main()
