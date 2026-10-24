# -*- coding: utf-8 -*-
"""M2 台账生成器：Ledger → 《比选输出模板》5-sheet xlsx。

模板优先：存在《比选输出模板.xlsx》时在副本上填写（保留企业原格式，零迁移成本）；
否则按模板表头新建（GitHub 独立检出场景）。
表头与模板逐字一致；不自动推荐最低价、不裁决异常，只呈现计算结果与线索（规则第五、七条）。
"""
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# 供应商清单完全由数据决定（每次比选 3–6 家），不从常量来；
# 对比表的供应商列名 = f"{供应商}含税金额"，写模板时按表头匹配、缺列自动补插。

SHEET_ORDER = ["对比表", "需补充事项", "操作记录", "报价明细", "资格材料清单"]

FIXED_HEADERS = {
    "需补充事项": ["事项编号", "供应商", "清单编号", "问题", "所需材料", "处理状态", "依据"],
    "操作记录": ["操作时间", "操作者", "供应商", "分项", "原值", "新值", "变更依据"],
    "报价明细": ["供应商", "清单编号", "原始分项", "原单位", "原数量", "原单价", "原税率",
               "原金额", "统一单位", "统一数量", "未税金额", "税额", "含税金额", "折算依据", "来源位置"],
    "资格材料清单": ["供应商", "材料类别", "文件编号", "提交情况", "有效期", "待补充事项", "复核状态"],
}
DB_FIXED = ["清单编号", "名称", "单位", "数量"]
DB_TAIL = ["差异及待核对", "来源定位"]


def _headers_for(name, sups):
    if name == "对比表":
        return DB_FIXED + [f"{s}含税金额" for s in sups] + DB_TAIL
    return FIXED_HEADERS[name]


headers_for = _headers_for   # 对外导出名

WIDTHS = {
    "对比表": [10, 18, 6, 8, 14, 14, 14, 14, 42, 46],
    "需补充事项": [11, 7, 9, 38, 26, 16, 42],
    "操作记录": [17, 10, 7, 8, 20, 20, 48],
    "报价明细": [7, 9, 17, 7, 9, 10, 9, 11, 9, 9, 12, 10, 12, 44, 34],
    "资格材料清单": [7, 14, 12, 9, 12, 32, 22],
}

NUM_FMT = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOP = Alignment(vertical="top", wrap_text=True)

# 对比表"差异及待核对"列里不重复展示的旗标（单元格本身已体现，或属纯信息性）
EXCLUDE_DIFF = ("税率未给", "单价未填", "报价中无此分项", "未知值不当0", "按列示金额")

# 需补充事项：旗标/备注关键词 → (问题, 所需材料)
PROBLEMS = {
    "税率": ("税率未声明，含税口径不可折算（未默认补写）", "税率与开票口径书面确认"),
    "漏报": ("该分项漏报/未报价", "补充该分项单价与金额"),
    "重复": ("同一分项重复报价", "确认重复行，只计一次"),
    "数量": ("报价数量与统一清单不符", "更正数量并说明依据"),
    "抄写": ("列示金额与数量×单价乘积不符", "按乘积更正分项金额"),
    "单位": ("报价单位与清单单位不一致（如项≠台班）", "单位换算的供应商书面说明"),
    "附件": ("费用附件未并入主报价", "确认附件并入口径"),
    "合计": ("原报列示合计与明细加总不符", "确认明细后重算合计"),
    "区间": ("单价超出参考价格区间（异常线索）", "报价依据说明（范围理解/笔误/其他）"),
}


def _cat_of(text: str):
    if "税率未给" in text:
        return "税率"
    if ("无此分项" in text) or ("单价未填" in text):
        return "漏报"
    if "重复" in text:
        return "重复"
    if "数量与清单不符" in text:
        return "数量"
    if "列示金额与数量×单价差" in text:
        return "抄写"
    if "单位口径待澄清" in text:
        return "单位"
    if "未并入" in text:
        return "附件"
    if "明细加总" in text:
        return "合计"
    if "参考区间" in text:
        return "区间"
    return None


def _closing_audit(led, sup, item, cat):
    """澄清闭环判断：只认依据为 CLARIFY-01 的留痕（CHANGE-01 工程量变更不得关闭异常事项）。"""
    for a in led.audit:
        if a["供应商"] != sup or "CLARIFY-01" not in a["变更依据"]:
            continue
        if cat == "税率":
            if a["分项"] == "税率":
                return a
        elif cat == "合计":
            return a
        elif a["分项"] == item:
            return a
    return None


def _need_items(led, qualifiers=None, closure_led=None, invoice_status=None):
    """从旗标生成需补充事项（去重、带闭环状态）；资格缺件/疑点与发票类型未说明一并入表且保持开放。

    led: 问题发现来源台账（传 V1 原始版使清单跨版本累计）
    closure_led: 闭环求值台账（当前版本；缺省同 led）
    invoice_status: {供应商: {"declared": bool, "invoice_type": str|None}} 发票类型影响真实成本口径
    """
    closure_led = closure_led or led
    items, seen = [], set()

    def add(sup, item, cat, evidence, closed=None, problem=None, material=None, basis=None):
        # 税率/合计类按供应商折叠（逐行旗标不重复立事项）
        row_item = "—" if cat in ("税率", "合计") else item
        key = (sup, row_item, cat)
        if key in seen:
            return
        seen.add(key)
        if closed is None:
            a = _closing_audit(closure_led, sup, row_item, cat)
            closed, basis = a is not None, (f"CLARIFY-01：{a['原值']} → {a['新值']}" if a else evidence)
        items.append({
            "supplier": sup, "item": row_item, "cat": cat,
            "problem": problem or PROBLEMS[cat][0], "material": material or PROBLEMS[cat][1],
            "closed": closed, "basis": basis,
        })

    for item, per in led.rows.items():
        for sup in per:
            for c in per[sup]:
                # 附件单列场景（如丙C13）已由"未并入"事项覆盖，不再另立泛化漏报项
                has_attach = any("未并入" in f for f in c.flags)
                for f in c.flags:
                    cat = _cat_of(f)
                    if cat == "漏报" and has_attach:
                        continue
                    if cat:
                        add(sup, item, cat, c.source)
    for sup in led.totals:
        for f in led.totals[sup]["flags"]:
            cat = _cat_of(f)
            if cat:
                add(sup, "—", cat, "原报价文件列示合计")
    for sup, st in (invoice_status or {}).items():
        if st.get("declared") and not st.get("invoice_type"):
            add(sup, "税率", "发票",
                "税率已声明；发票类型未说明",
                closed=False,
                problem="发票类型未说明（专票可抵扣则比价看未税，普票含税即真实成本）",
                material="发票类型说明（专票/普票）",
                basis="税率已声明；发票类型未说明，双口径并排供商务判断")
    for q in (qualifiers or []):
        if q.get("submitted") != "已提交" or "疑点" in q.get("review", "") \
                or "缺件" in q.get("review", ""):
            items.append({
                "supplier": q["supplier"], "item": "资质", "cat": "资质",
                "problem": q.get("note") or "资格材料待补充",
                "material": "补充/续期资质材料",
                "closed": False,  # 资格事项保持开放至评审确认（CHANGE-01：待补事项不关闭）
                "basis": "资格材料登记.xlsx；系统只整理缺件与疑点，有效性由评审人员判断",
            })
    return items


def _amount_or_text(c, value):
    if c.status != "已计算":
        return "未报"
    return "税率未给" if value is None else float(value)


def _diff_text(per, sups):
    parts = []
    for sup in sups:
        fs = [f for c in per[sup] for f in c.flags
              if not any(k in f for k in EXCLUDE_DIFF)]
        if fs:
            parts.append(f"{sup}：" + "；".join(fs))
    return "；".join(parts)


def _src_text(per, sups):
    parts = [f"{sup}：{per[sup][0].source}" for sup in sups
             if per[sup][0].source != "—"]
    return "；".join(parts)


def _put(ws, r, col, v, num=False, bold=False):
    cell = ws.cell(row=r, column=col, value=float(v) if num and isinstance(v, Decimal) else v)
    cell.alignment = TOP
    if num:
        cell.number_format = NUM_FMT
    if bold:
        cell.font = Font(bold=True)
    return cell


def _fill_headers(sups, template_path):
    if template_path and Path(template_path).exists():
        wb = load_workbook(template_path)
        for name in SHEET_ORDER:
            if name not in wb.sheetnames:
                ws = wb.create_sheet(name)
                for i, h in enumerate(_headers_for(name, sups), 1):
                    ws.cell(row=1, column=i, value=h)
        for name in SHEET_ORDER:
            ws = wb[name]
            if ws.max_row and ws.max_row > 1:
                ws.delete_rows(2, ws.max_row)
    else:
        wb = Workbook()
        wb.remove(wb.active)
        for name in SHEET_ORDER:
            ws = wb.create_sheet(name)
            for i, h in enumerate(_headers_for(name, sups), 1):
                ws.cell(row=1, column=i, value=h)
    return wb


def _db_colmap(ws, sups):
    """对比表列定位：按表头匹配「{供应商}含税金额」；模板缺该供应商列时在「差异及待核对」前补插。"""
    def scan():
        return {ws.cell(row=1, column=c).value: c
                for c in range(1, ws.max_column + 1) if ws.cell(row=1, column=c).value}
    hdr = scan()
    for s in sups:
        if not any(isinstance(h, str) and h.startswith(s) and h.endswith("含税金额")
                   for h in hdr):
            diff_col = hdr["差异及待核对"]
            ws.insert_cols(diff_col)
            ws.cell(row=1, column=diff_col, value=f"{s}含税金额")
            hdr = scan()
    return hdr


def _style(wb, sups):
    for name in SHEET_ORDER:
        ws = wb[name]
        headers = _headers_for(name, sups)
        if name == "对比表":
            widths = [10, 18, 6, 8] + [14] * len(sups) + [42, 46]
        else:
            widths = WIDTHS[name]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        for i in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=i)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = TOP
        ws.freeze_panes = "A2"


def write_ledger(led, catalog, out_path, rate_map, qualifiers=None,
                 qty_map=None, prev_totals=None, version_note="", template_path=None,
                 issue_source=None, invoice_status=None):
    """把一个版本的 Ledger 写成 5-sheet 台账 xlsx。

    rate_map: {供应商: "6%"/"未声明"} 原税率列展示
    qty_map:  对比表"数量"列（统一清单数量，变更后为调整值）
    prev_totals: 上一版本合计（V2 写"较V1金额差异"行，落实变更通知的"保留基线与金额差异"）
    issue_source: 需补充事项的来源台账（传 V1 原始版，使问题清单跨版本累计、
                  处理状态再对当前版本求值；缺省用 led 自身）
    invoice_status: {供应商: {"declared","invoice_type"}} 发票类型未说明时立事项（M5）
    """
    qty_map = qty_map or {}
    sups = list(led.totals.keys())
    wb = _fill_headers(sups, template_path)

    # ---------------- 对比表 ----------------
    ws = wb["对比表"]
    colmap = _db_colmap(ws, sups)
    r = 2
    for item, per in led.rows.items():
        cat = catalog[item]
        _put(ws, r, colmap["清单编号"], item)
        _put(ws, r, colmap["名称"], cat["name"])
        _put(ws, r, colmap["单位"], cat["unit"])
        _put(ws, r, colmap["数量"], float(qty_map.get(item, cat["qty"])))
        for sup in sups:
            ls = per[sup]
            if ls[0].status != "已计算":
                v = "未报"
            elif ls[0].hs is None:
                v = "税率未给"
            else:
                v = sum(x.hs for x in ls)   # 同分项多行（重复未澄清）按列示原样合计
            _put(ws, r, colmap[f"{sup}含税金额"], v, num=not isinstance(v, str))
        _put(ws, r, colmap["差异及待核对"], _diff_text(per, sups))
        _put(ws, r, colmap["来源定位"], _src_text(per, sups))
        r += 1

    hs = lambda sup: led.totals[sup]["hs"]
    _put(ws, r, colmap["清单编号"], "合计", bold=True)
    _put(ws, r, colmap["名称"], "按明细计算", bold=True)
    for sup in sups:
        v = hs(sup)
        _put(ws, r, colmap[f"{sup}含税金额"], v if v is not None else "—",
             num=v is not None, bold=True)
    _put(ws, r, colmap["差异及待核对"],
         "；".join(f"{s}：{f}" for s in sups for f in led.totals[s]["flags"]))
    r += 1
    _put(ws, r, colmap["清单编号"], "列示合计")
    _put(ws, r, colmap["名称"], "供应商原报，未核实")
    for sup in sups:
        st = led.totals[sup]["stated"]
        _put(ws, r, colmap[f"{sup}含税金额"], st, num=st is not None)
    r += 1
    if prev_totals is not None:
        _put(ws, r, colmap["清单编号"], "较V1差异")
        _put(ws, r, colmap["名称"], "V2合计−V1澄清后合计（含税口径）")
        for sup in sups:
            if hs(sup) is not None and prev_totals.get(sup) is not None:
                d = hs(sup) - prev_totals[sup]
                _put(ws, r, colmap[f"{sup}含税金额"], f"{float(d):+,.2f}")
        r += 1
    r += 1
    for note in (version_note, "金额单位：元；本表由规则引擎生成，金额计算零LLM参与；异常仅提示，价格合理性与选商由评审人员判断"):
        if note:
            _put(ws, r, colmap["清单编号"], note)
            r += 1

    # ---------------- 需补充事项 ----------------
    ws = wb["需补充事项"]
    r = 2
    for i, it in enumerate(_need_items(issue_source or led, qualifiers,
                                       closure_led=led, invoice_status=invoice_status), 1):
        _put(ws, r, 1, f"QB-{i:03d}")
        _put(ws, r, 2, it["supplier"])
        _put(ws, r, 3, it["item"])
        _put(ws, r, 4, it["problem"])
        _put(ws, r, 5, it["material"])
        _put(ws, r, 6, "已澄清关闭" if it["closed"] else "待补充")
        _put(ws, r, 7, it["basis"])
        r += 1

    # ---------------- 操作记录 ----------------
    ws = wb["操作记录"]
    r = 2
    for a in led.audit:
        for c, key in enumerate(("操作时间", "操作者", "供应商", "分项", "原值", "新值", "变更依据"), 1):
            _put(ws, r, c, a[key])
        r += 1

    # ---------------- 报价明细 ----------------
    ws = wb["报价明细"]
    r = 2
    for sup in sups:
        for item, per in led.rows.items():
            cat = catalog[item]
            for c in per[sup]:
                _put(ws, r, 1, sup)
                _put(ws, r, 2, item)
                _put(ws, r, 3, c.name)
                _put(ws, r, 4, c.unit)
                _put(ws, r, 5, float(c.orig_qty) if c.orig_qty is not None else "—")
                _put(ws, r, 6, c.unit_price, num=c.unit_price is not None)
                _put(ws, r, 7, rate_map.get(sup, "—"))
                _put(ws, r, 8, c.stated_amount, num=c.stated_amount is not None)
                _put(ws, r, 9, cat["unit"])
                _put(ws, r, 10, float(c.qty) if c.qty is not None else "—")
                for col, val in ((11, c.ws), (12, c.tax_amt), (13, c.hs)):
                    v = _amount_or_text(c, val)
                    _put(ws, r, col, v, num=not isinstance(v, str))
                _put(ws, r, 14, c.basis)
                _put(ws, r, 15, c.source)
                r += 1

    # ---------------- 资格材料清单 ----------------
    ws = wb["资格材料清单"]
    r = 2
    for q in (qualifiers or []):
        _put(ws, r, 1, q["supplier"])
        _put(ws, r, 2, q["category"])
        _put(ws, r, 3, q["doc"] or "—")
        _put(ws, r, 4, q["submitted"])
        _put(ws, r, 5, q["validity"] or "—")
        _put(ws, r, 6, q["note"] or "—")
        _put(ws, r, 7, q["review"])
        r += 1

    _style(wb, sups)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
