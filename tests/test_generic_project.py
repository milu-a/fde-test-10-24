# -*- coding: utf-8 -*-
"""通用性证明：与 GC-017 无关的假项目——3 行清单、2 家供应商、不同单位与税率，
直接以数据构造走完 引擎计算 → 台账生成 → 逻辑变更重算，验证"标准比价清单是啥都支持"。
"""
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook, load_workbook

from engine.ledger import LedgerBuilder
from engine.models import Direction, QuoteLine, SourceRef, TaxProfile
from engine.money import D
from report import SHEET_ORDER, write_ledger
from report.xlsx import _db_colmap

CATALOG = {
    "X1": {"id": "X1", "name": "场地平整", "unit": "㎡", "qty": D("500"), "range": ["3", "6"]},
    "X2": {"id": "X2", "name": "界桩埋设", "unit": "个", "qty": D("20"), "range": ["40", "60"]},
    "X3": {"id": "X3", "name": "竣工测量", "unit": "幅", "qty": D("2"), "range": ["800", "1200"]},
}


def _line(sup, item, unit, qty, price, stated, direction):
    return QuoteLine(supplier=sup, item_id=item, name=CATALOG[item]["name"], unit=unit,
                     qty=D(qty), unit_price=D(price) if price else None,
                     direction=direction, stated_amount=D(stated) if stated else None,
                     source=SourceRef(f"{sup}报价.xlsx", f"条目{item}"))


def _make_builder():
    tax_wu = TaxProfile("戊")
    tax_wu.declare("0.03", SourceRef("戊报价.xlsx", "税率列"), "2026-01-01", "戊")
    suppliers = {
        "戊": {"direction": Direction.TAX_EXCL,
              "lines": [_line("戊", "X1", "㎡", "500", "4.00", "2000.00", Direction.TAX_EXCL),
                        _line("戊", "X2", "个", "20", "50.00", "1000.00", Direction.TAX_EXCL)],
              "stated_total": D("3000.00")},
        "己": {"direction": Direction.TAX_INCL,
              "lines": [_line("己", "X1", "㎡", "500", "4.30", "2150.00", Direction.TAX_INCL)],
              "stated_total": D("2150.00")},
    }
    return LedgerBuilder(CATALOG, suppliers, {"戊": tax_wu, "己": TaxProfile("己")})


class GenericProjectTest(unittest.TestCase):
    def test_engine_generic(self):
        led = _make_builder().build("测试项目V1")
        self.assertEqual(set(led.totals), {"戊", "己"})
        self.assertEqual(led.totals["戊"]["ws"], D("3000.00"))   # 未税方向合计
        self.assertEqual(led.totals["戊"]["hs"], D("3090.00"))   # +ROUND(3000×3%)
        self.assertEqual(led.totals["己"]["hs"], D("2150.00"))   # 含税方向合计
        self.assertIsNone(led.totals["己"]["ws"])                # 己税率未声明
        self.assertEqual(led.rows["X2"]["己"][0].status, "未报")  # 漏报占位
        # 任意分项的工程量变更照常重算
        v2 = _make_builder().build("测试项目V2", qty_overrides={"X1": "600"})
        self.assertEqual(v2.rows["X1"]["戊"][0].ws, D("2400.00"))

    def test_report_dynamic_suppliers(self):
        led = _make_builder().build("测试项目V1")
        with tempfile.TemporaryDirectory() as td:
            p = write_ledger(led, CATALOG, Path(td) / "x.xlsx",
                             {"戊": "3%", "己": "未声明"})
            wb = load_workbook(p)
        self.assertEqual(wb.sheetnames, SHEET_ORDER)
        ws = wb["对比表"]
        got = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        self.assertEqual(got, ["清单编号", "名称", "单位", "数量",
                              "戊含税金额", "己含税金额", "差异及待核对", "来源定位"])
        r_x1 = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "X1")
        self.assertEqual(ws.cell(row=r_x1, column=5).value, 2060.0)   # 戊X1：2000+ROUND(2000×3%)=2060
        self.assertEqual(ws.cell(row=r_x1, column=6).value, 2150.0)   # 己X1含税
        r_x2 = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "X2")
        self.assertEqual(ws.cell(row=r_x2, column=5).value, 1030.0)   # 戊X2
        self.assertEqual(ws.cell(row=r_x2, column=6).value, "未报")

    def test_template_missing_supplier_column_inserted(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "对比表"
        for i, h in enumerate(["清单编号", "名称", "单位", "数量", "戊含税金额",
                               "差异及待核对", "来源定位"], 1):
            ws.cell(row=1, column=i, value=h)
        colmap = _db_colmap(ws, ["戊", "己"])   # 模板没有己 → 自动补插
        self.assertEqual(colmap["戊含税金额"], 5)
        self.assertEqual(colmap["己含税金额"], 6)
        self.assertEqual(colmap["差异及待核对"], 7)
        self.assertEqual(ws.cell(row=1, column=6).value, "己含税金额")


if __name__ == "__main__":
    unittest.main()
