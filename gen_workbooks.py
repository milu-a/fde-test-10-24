# -*- coding: utf-8 -*-
"""M2 台账生成器运行入口：GC-017 三阶段 → 三份 5-sheet 比选台账 xlsx。

运行：python gen_workbooks.py   （在 agent/ 目录下）
"""
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

from engine import load_all
from engine.money import fmt
from report import write_ledger

TEMPLATE = ROOT.parent / "04 报价各说各话，工程分包怎么比？" / "01 原始报价与清单" / "比选输出模板.xlsx"
OUT = ROOT / "out"


def main():
    ns = load_all()
    b = ns.builder
    ch = ns.events["change"]
    catalog = ns.catalog
    qualifiers = json.loads((ROOT / "data" / "gc017_qualifications.json").read_text("utf-8"))["rows"]

    v1 = b.build("V1 原始报价（澄清前）")
    v1c = b.build("V1 澄清后（CLARIFY-01 已回填）", amendments=ns.events["clarify"])
    v2 = b.build("V2 变更后（CHANGE-01：C08 60→72根，V1基线保留）",
                 amendments=ns.events["clarify"], qty_overrides=ch["qty_overrides"],
                 qty_basis=ch["id"], qty_at=ch["at"], qty_operator=ch["operator"])

    catalog_qty = {i: c["qty"] for i, c in catalog.items()}
    qty_v2 = {**catalog_qty, "C08": catalog["C08"]["qty"] + 12}  # 60→72

    stages = [
        ("GC-017_比选台账_V1原始.xlsx", v1,
         {"甲": "6%", "乙": "未声明", "丙": "6%", "丁": "6%"},
         catalog_qty, None,
         "台账版本：V1 原始（澄清前，异常线索最全；乙税率未声明，含税口径不可比）"),
        ("GC-017_比选台账_V1澄清后.xlsx", v1c,
         {"甲": "6%", "乙": "6%", "丙": "6%", "丁": "6%"},
         catalog_qty, None,
         "台账版本：V1 澄清后（CLARIFY-01 已回填，全程留痕见「操作记录」）"),
        ("GC-017_比选台账_V2变更后.xlsx", v2,
         {"甲": "6%", "乙": "6%", "丙": "6%", "丁": "6%"},
         qty_v2, {s: v1c.totals[s]["hs"] for s in ("甲", "乙", "丙", "丁")},
         "台账版本：V2 变更后（CHANGE-01：C08低应变 60根→72根，四家单价税率沿用；V1基线与V1澄清后台账另行保留）"),
    ]

    print("生成比选台账（M2）")
    print("=" * 60)
    for fname, led, rate_map, qmap, prev, note in stages:
        tpl = TEMPLATE if TEMPLATE.exists() else None
        if tpl is None and not getattr(main, "_warned", False):
            print("（未找到企业模板文件，按模板表头新建工作簿）")
            main._warned = True
        p = write_ledger(led, catalog, OUT / fname, rate_map,
                         qualifiers=qualifiers, qty_map=qmap,
                         prev_totals=prev, version_note=note, template_path=tpl,
                         issue_source=v1)
        from openpyxl import load_workbook
        wb = load_workbook(p)
        counts = {s: max(wb[s].max_row - 1, 0) for s in wb.sheetnames}
        print(f"  {p.name}")
        print(f"    sheets: {' / '.join(wb.sheetnames)}")
        print(f"    行数: " + "，".join(f"{s}{n}" for s, n in counts.items()))

    print(f"\n输出目录：{OUT}")
    print("提示：三版台账互不覆盖——V1 原始与 V1 澄清后保留基线（规则第六条：原件永不覆盖）。")


if __name__ == "__main__":
    main()
