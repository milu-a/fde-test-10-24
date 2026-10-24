#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
【纯代码】读取器 —— Excel 与 PDF 文字层
==========================================
原理：这两类文件内部"本来就有字"（结构化数据/字符指令），
      按文件格式规范直接解析，无AI、无token、确定性100%。

对外提供：
  read_excel(path)            → (records, note)          Excel通道
  read_pdf_text(path)         → (records, note) | None   PDF文字层通道（None=无文字层）
  pdf_has_text(path)          → bool
  structure_text_document(lines, src) → (records, note)  文字行→结构化记录
      （供模型读取器复用：模型只负责把像素变成文字行，
        把文字行变成记录这一步仍是纯代码）
"""
import re
from pathlib import Path

import openpyxl

FIELDS = ["行类型", "原始编码", "原始分项名称", "单位", "数量", "单价", "单价口径",
          "税率", "税率依据", "原报金额", "置信度", "待人工标记", "出处-文件", "出处-定位", "备注"]

UNIT_TOKENS = "点|亩|孔|根|构件|组|项|台班"
ROW_RE = re.compile(r"^(C\d{2})\s+(.+?)\s+(" + UNIT_TOKENS + r")\s+(\d+)(?:\s+([\d.]+))?(?:\s+([\d.]+))?\s*$")
HEADER_KEYS = ("分项编号", "编码", "清单编号")
META_KEYS = ("项目编号", "报价单位", "报价版本", "版本", "费用范围", "报价日期", "联系人")
TOTAL_KEYS = ("合计", "总计", "小计")


def _rec(d: dict) -> dict:
    r = {k: "" for k in FIELDS}
    r.update(d)
    return r


# ---------------------------------------------------------------- Excel
def read_excel(path: Path):
    """xlsx 本质是ZIP+XML：openpyxl 按格式规范读单元格，坐标原生可得。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    target = None
    for ws in wb.worksheets:
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            cells = [str(c).strip() if c is not None else "" for c in r]
            if any(k in cells for k in HEADER_KEYS) and "单位" in cells:
                target = (ws, i + 1, cells)
                break
        if target:
            break
    if not target:
        raise ValueError("未找到清单表头（分项编号/编码 + 单位）")

    ws, header_row_no, cells = target
    header = {h: j for j, h in enumerate(cells) if h}

    def col(*keys, default=None):
        for k in keys:
            for h, j in header.items():
                if k in h:
                    return j
        return default

    c_code = col("分项编号", "编码", default=0)
    c_name = col("名称", "分项", default=1)
    c_unit = col("单位", default=2)
    c_qty = col("数量", default=3)
    c_price = next((j for h, j in header.items() if "单价" in h), None)
    c_tax = next((j for h, j in header.items() if "税率" in h), None)
    c_amt = next((j for h, j in header.items() if "金额" in h), None)
    price_hdr = next((h for h in header if "单价" in h), "")
    basis = "含税" if "含税" in price_hdr else ("未税" if "未税" in price_hdr else "未声明")

    version, meta_pairs = "", []
    for ws2 in wb.worksheets:
        for r in ws2.iter_rows(values_only=True):
            if r and r[0] is not None and str(r[0]).strip() in META_KEYS:
                val = r[1] if len(r) > 1 and r[1] is not None else ""
                meta_pairs.append(f"{str(r[0]).strip()}：{val}")
                if str(r[0]).strip() in ("报价版本", "版本"):
                    version = str(val).strip()

    records, n_items = [], 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        row_no = i + 1
        if row_no <= header_row_no or not r:
            continue
        code = str(r[c_code]).strip() if r[c_code] is not None else ""
        tax_val = r[c_tax] if c_tax is not None else None
        tax_disp = f"{float(tax_val):.0%}" if isinstance(tax_val, (int, float)) else "未给"
        price_v = r[c_price] if c_price is not None else None
        amt_v = r[c_amt] if c_amt is not None else None
        loc = f"工作表[{ws.title}] 第{row_no}行"
        if re.fullmatch(r"[A-Z]\d{2}", code):
            n_items += 1
            flag = ""
            if tax_disp == "未给":
                flag = "税率未提供-待商务确认（不默认补写）"
            if price_v is None:
                flag = (flag + "；" if flag else "") + "单价缺失-待澄清"
            records.append(_rec({
                "行类型": "分项", "原始编码": code,
                "原始分项名称": str(r[c_name]).strip() if r[c_name] else "",
                "单位": str(r[c_unit]).strip() if r[c_unit] else "",
                "数量": r[c_qty], "单价": price_v, "单价口径": basis, "税率": tax_disp,
                "税率依据": "行内税率列" if tax_disp != "未给" else "全文未给税率",
                "原报金额": amt_v, "置信度": "高", "待人工标记": flag,
                "出处-文件": path.name, "出处-定位": loc,
                "备注": f"版本{version}" if version else ""}))
        elif code in TOTAL_KEYS:
            records.append(_rec({
                "行类型": "合计", "原始分项名称": f"报价{code}", "单价口径": basis,
                "税率": tax_disp, "原报金额": amt_v, "置信度": "高",
                "出处-文件": path.name, "出处-定位": loc}))
    if meta_pairs:
        records.append(_rec({
            "行类型": "说明", "原始分项名称": "报价说明", "备注": "；".join(meta_pairs),
            "置信度": "高", "出处-文件": path.name, "出处-定位": "说明区"}))
    if n_items == 0:
        raise ValueError("找到表头但未解析到分项行（C01式编码）")
    return records, f"Excel通道（纯代码）：分项{n_items}行，口径={basis}"


# ------------------------------------------------------------- PDF 文字层
def pdf_has_text(path: Path) -> bool:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return any((p.extract_text() or "").strip() for p in pdf.pages)


def read_pdf_text(path: Path):
    """文字版PDF内部存着"在哪个坐标画哪个字"的指令：pdfplumber抽出字符按位置拼行。"""
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        pages = [(i + 1, p.extract_text() or "") for i, p in enumerate(pdf.pages)]
    if not any(t.strip() for _, t in pages):
        return None  # 无文字层 → 交给模型读取器
    records, n_items = [], 0
    version, basis_note = "", ""
    head = "\n".join(t for _, t in pages)[:400]
    mver = re.search(r"([A-Z]-V\d+)", head)
    if mver:
        version = mver.group(1)
    basis = "未税" if "未税" in head else ("含税" if "含税" in head else "未声明")
    mtax = re.search(r"税率\s*(?:为|是|：|:)?\s*(\d+(?:\.\d+)?)\s*%", head)
    tax_disp = f"{int(float(mtax.group(1)))}%" if mtax else "未给"
    tax_src = "报价页首说明" if mtax else "全文未给税率"
    for pno, text in pages:
        for idx, line in enumerate([l.strip() for l in text.splitlines() if l.strip()], 1):
            loc = f"第{pno}页 文本第{idx}行"
            m = ROW_RE.match(line)
            if m:
                n_items += 1
                code, name, unit, qty, price, amount = m.groups()
                price_v = float(price) if price else None
                flag = ""
                if price_v is None:
                    flag = "单价缺失（漏报）-待澄清补充"
                if tax_disp == "未给":
                    flag = (flag + "；" if flag else "") + "税率未提供-待商务确认（不默认补写）"
                records.append(_rec({
                    "行类型": "分项", "原始编码": code, "原始分项名称": name,
                    "单位": unit, "数量": int(qty), "单价": price_v, "单价口径": basis,
                    "税率": tax_disp, "税率依据": tax_src,
                    "原报金额": float(amount) if amount else None, "置信度": "高",
                    "待人工标记": flag, "出处-文件": path.name, "出处-定位": loc,
                    "备注": f"版本{version}" if version else ""}))
                continue
            mt = re.search(r"合计[：:]([\d.]+)元?", line)
            if mt:
                records.append(_rec({
                    "行类型": "合计", "原始分项名称": f"报价列示合计（{basis}口径）",
                    "单价口径": basis, "税率": tax_disp, "税率依据": tax_src,
                    "原报金额": float(mt.group(1)), "置信度": "高",
                    "出处-文件": path.name, "出处-定位": loc}))
                continue
            if any(k in line for k in ("费用说明", "联系人", "口径")) and len(line) < 120:
                basis_note = (basis_note + "；" + line).strip("；")
    if basis_note:
        records.append(_rec({
            "行类型": "说明", "原始分项名称": "报价说明", "单价口径": basis,
            "税率": tax_disp, "税率依据": tax_src, "备注": basis_note, "置信度": "高",
            "出处-文件": path.name, "出处-定位": "说明行"}))
    if n_items == 0:
        return None  # 有文字但不像清单 → 交给人工
    return records, f"PDF文字层通道（纯代码）：分项{n_items}行，口径={basis}，税率={tax_disp}"


# ------------------------------------------- 文字行 → 结构化记录（共用）
def structure_text_document(lines, src_label: str):
    """
    把"整份文件的文字行"（PDF文字层或OCR输出的文字行）变成结构化记录。
    模型读取器复用此函数：模型只负责像素→文字行，结构化仍是纯代码。
    """
    records, n_items = [], 0
    head = "\n".join(lines)[:400]
    mver = re.search(r"([A-Z]-V\d+)", head)
    version = mver.group(1) if mver else ""
    basis = "未税" if "未税" in head else ("含税" if "含税" in head else "未声明")
    mtax = re.search(r"税率\s*(?:为|是|：|:)?\s*(\d+(?:\.\d+)?)\s*%", head)
    tax_disp = f"{int(float(mtax.group(1)))}%" if mtax else "未给"
    tax_src = "原文说明" if mtax else "全文未给税率"
    note_lines = []
    ocr_flag = "OCR数字待人工抽检"
    for idx, line in enumerate(lines, 1):
        loc = f"第{idx}文字行"
        m = ROW_RE.match(line)
        if m:
            n_items += 1
            code, name, unit, qty, price, amount = m.groups()
            price_v = float(price) if price else None
            flag = ""
            if price_v is None:
                flag = "单价缺失（漏报）-待澄清补充"
            if tax_disp == "未给":
                flag = (flag + "；" if flag else "") + "税率未提供-待商务确认（不默认补写）"
            records.append(_rec({
                "行类型": "分项", "原始编码": code, "原始分项名称": name, "单位": unit,
                "数量": int(qty), "单价": price_v, "单价口径": basis, "税率": tax_disp,
                "税率依据": tax_src, "原报金额": float(amount) if amount else None,
                "置信度": "中（模型识别）",
                "待人工标记": (flag + "；" + ocr_flag) if flag else ocr_flag,
                "出处-文件": src_label, "出处-定位": loc,
                "备注": f"版本{version}" if version else ""}))
            continue
        mt = re.search(r"合计[：:]?\s*([\d.]+)元?", line)
        if mt:
            records.append(_rec({
                "行类型": "合计", "原始分项名称": f"报价合计（{basis}口径）",
                "单价口径": basis, "税率": tax_disp, "税率依据": tax_src,
                "原报金额": float(mt.group(1)), "置信度": "中（模型识别）",
                "出处-文件": src_label, "出处-定位": loc}))
        elif any(k in line for k in ("费用说明", "联系人", "口径", "报价")) and len(line) < 120:
            note_lines.append(line)
    if note_lines:
        records.append(_rec({
            "行类型": "说明", "原始分项名称": "报价说明", "单价口径": basis,
            "税率": tax_disp, "税率依据": tax_src, "备注": "；".join(note_lines[:5]),
            "置信度": "中（模型识别）", "出处-文件": src_label, "出处-定位": "说明行"}))
    return records, f"文字行结构化：分项{n_items}行，口径={basis}，税率={tax_disp}"
